from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


file_name = r"C:\Users\S\OneDrive\Desktop\ملکیان\nova\tables_description.xlsx"
wb = Workbook()

# حذف Sheet پیش فرض
default = wb.active
wb.remove(default)


tables = {

"customers": [
("customer_id","INT","PK","شناسه یکتا مشتری"),
("first_name","VARCHAR","-","نام مشتری"),
("last_name","VARCHAR","-","نام خانوادگی مشتری"),
("email","VARCHAR","-","ایمیل مشتری"),
("phone","VARCHAR","-","شماره تماس"),
("birth_date","DATE","-","تاریخ تولد"),
("city","VARCHAR","-","شهر مشتری"),
("created_date","DATE","-","تاریخ ثبت مشتری")
],


"products": [
("product_id","INT","PK","شناسه محصول"),
("product_name","VARCHAR","-","نام محصول"),
("category","VARCHAR","-","دسته بندی محصول"),
("brand","VARCHAR","-","برند محصول"),
("price","FLOAT","-","قیمت فروش"),
("cost","FLOAT","-","هزینه خرید"),
("status","VARCHAR","-","وضعیت محصول")
],


"sales": [
("sale_id","INT","PK","شناسه فروش"),
("customer_id","INT","FK","شناسه مشتری"),
("product_id","INT","FK","شناسه محصول"),
("branch_id","INT","FK","شناسه شعبه"),
("sale_date","DATE","-","تاریخ فروش"),
("quantity","INT","-","تعداد محصول"),
("total_amount","FLOAT","-","مبلغ کل فروش")
],


"orders": [
("order_id","INT","PK","شناسه سفارش"),
("customer_id","INT","FK","شناسه مشتری"),
("order_date","DATE","-","تاریخ سفارش"),
("status","VARCHAR","-","وضعیت سفارش"),
("total_amount","FLOAT","-","مبلغ سفارش")
],


"order_details": [
("detail_id","INT","PK","شناسه جزئیات سفارش"),
("order_id","INT","FK","شناسه سفارش"),
("product_id","INT","FK","شناسه محصول"),
("quantity","INT","-","تعداد"),
("unit_price","FLOAT","-","قیمت واحد")
],


"payments": [
("payment_id","INT","PK","شناسه پرداخت"),
("sale_id","INT","FK","شناسه فروش"),
("payment_method","VARCHAR","-","روش پرداخت"),
("amount","FLOAT","-","مبلغ پرداخت"),
("payment_date","DATE","-","تاریخ پرداخت")
],


"inventory": [
("inventory_id","INT","PK","شناسه موجودی"),
("product_id","INT","FK","شناسه محصول"),
("stock_quantity","INT","-","تعداد موجودی"),
("reorder_level","INT","-","حد سفارش مجدد"),
("update_date","DATE","-","تاریخ بروزرسانی")
],


"branches": [
("branch_id","INT","PK","شناسه شعبه"),
("branch_name","VARCHAR","-","نام شعبه"),
("city","VARCHAR","-","شهر"),
("opening_date","DATE","-","تاریخ افتتاح")
],


"marketing_campaigns": [
("campaign_id","INT","PK","شناسه کمپین"),
("campaign_name","VARCHAR","-","نام کمپین"),
("channel","VARCHAR","-","کانال تبلیغات"),
("budget","FLOAT","-","بودجه"),
("start_date","DATE","-","تاریخ شروع"),
("end_date","DATE","-","تاریخ پایان")
],


"customer_feedback": [
("feedback_id","INT","PK","شناسه بازخورد"),
("customer_id","INT","FK","شناسه مشتری"),
("rating","INT","-","امتیاز مشتری"),
("customer_comment","TEXT","-","نظر مشتری"),
("feedback_date","DATE","-","تاریخ ثبت بازخورد")
]

}



for table_name, columns in tables.items():

    ws = wb.create_sheet(table_name)

    ws.append([
        "نام ستون",
        "نوع داده",
        "کلید",
        "توضیحات"
    ])


    for row in columns:
        ws.append(row)


    # طراحی Header
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(
            "solid",
            fgColor="D9EAF7"
        )
        cell.alignment = Alignment(
            horizontal="center"
        )


    # تنظیم عرض ستون ها
    widths = [20,15,15,45]

    for i,w in enumerate(widths,1):
        ws.column_dimensions[
            chr(64+i)
        ].width = w


    # Border
    for row in ws.iter_rows():
        for cell in row:
            cell.border = Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin")
            )

import os

wb.save(file_name)

print("tables_description.xlsx ساخته شد ✅")
print("مسیر فایل:")
print(os.path.abspath(file_name))
